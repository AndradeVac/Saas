"""Professional exports: sales report (Excel/PDF) and order/customer lists in Excel."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

EVIL = '=HYPERLINK("http://evil.example/x","Clique")'


def _order(shop, name="Maria Cliente"):
    product = shop.add_product(name="X-Bacon", price="30.00")
    response = shop.public_order(product["id"], customer_name=name, quantity=2, service_type="TAKEAWAY", payment_method="PIX")
    assert response.status_code == 201, response.text


def _book(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(response.content))


def test_sales_report_excel_has_every_sheet_with_real_numbers(make_tenant):
    shop = make_tenant()
    _order(shop)
    book = _book(shop.get("/analytics/dashboard/export", params={"format": "xlsx", "period": "month"}))
    assert book.sheetnames == ["Resumo", "Vendas por dia", "Horários", "Produtos", "Pedidos"]
    summary = book["Resumo"]
    assert summary["A6"].value == "Faturamento" and summary["A7"].value == 60.0
    products = [row for row in book["Produtos"].iter_rows(values_only=True) if row[1] == "X-Bacon"]
    assert products and products[0][2] == 2 and products[0][3] == 60.0
    assert book["Vendas por dia"]._charts  # native Excel chart


def test_sales_report_pdf_and_custom_dates(make_tenant):
    shop = make_tenant()
    _order(shop)
    pdf = shop.get("/analytics/dashboard/export", params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and len(pdf.content) > 2000
    long_ago = date.today() - timedelta(days=400)
    old = _book(shop.get("/analytics/dashboard/export", params={"format": "xlsx", "start": long_ago.isoformat(), "end": (long_ago + timedelta(days=1)).isoformat()}))
    assert old["Resumo"]["B7"].value == 0  # no orders in that window


def test_excel_lists_keep_customer_text_as_plain_text(make_tenant):
    shop = make_tenant()
    _order(shop, name=EVIL)
    for path in ("/orders/export", "/customers/export"):
        book = _book(shop.get(path, params={"format": "xlsx"}))
        cells = [cell for sheet in book for row in sheet.iter_rows() for cell in row]
        assert not any(cell.data_type == "f" for cell in cells), path
        assert any(cell.value == "'" + EVIL for cell in cells), path
