"""Professional exports: the sales report (Excel and PDF) and the order and customer lists (Excel).

Every file carries the restaurant's name, logo and brand colour, the period in words and the generation time.
Excel gets real numbers, dates and currency formats (sortable and summable), native charts, filters and frozen
headers; the PDF is an A4 report with headline numbers, a chart and ranked tables.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.labels import PAYMENT_LABELS, SERVICE_LABELS, STATUS_LABELS, WEEKDAYS, label
from app.core.spreadsheet import neutralise_formulas, safe_cell
from app.models.media import Media
from app.models.order import Order, OrderStatus, PaymentStatus
from app.models.tenant import Tenant
from app.repositories.order import OrderRepository
from app.schemas.analytics import DashboardAnalyticsResponse
from app.services.analytics import AnalyticsService
from app.services.media import MEDIA_PREFIX

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = '"R$" #,##0.00'
INK = "1F2933"
MUTED = "6B7280"
LINE = "E5E7EB"
ZEBRA = "F8F9FB"


def money(value) -> str:
    """R$ 1.234,56"""
    return "R$ " + f"{Decimal(value or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def percent(part, whole) -> float:
    return float(part) / float(whole) if whole else 0.0


def _hex(color: str | None) -> str:
    color = (color or "#c2410c").lstrip("#")
    return color.upper() if len(color) == 6 else "C2410C"


def _tint(hex_color: str, amount: float) -> str:
    """Mixes the brand colour with white (amount 0..1) for soft backgrounds."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(c + (255 - c) * amount):02X}" for c in (r, g, b))


def _logo_png(db: Session, tenant: Tenant) -> bytes | None:
    """The tenant's logo as PNG (stored as WebP in the database); None when missing or stored elsewhere."""
    url = tenant.logo_url or ""
    if not url.startswith(MEDIA_PREFIX):
        return None
    try:
        media = db.scalar(select(Media).where(Media.id == UUID(url.removeprefix(MEDIA_PREFIX)), Media.tenant_id == tenant.id))
        if media is None or not media.data:
            return None
        image = Image.open(io.BytesIO(media.data)).convert("RGBA")
        image.thumbnail((256, 256))
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    except Exception:  # a broken logo must never break the report
        return None


# ---- report data --------------------------------------------------------------------------------

@dataclass
class SalesReport:
    tenant: Tenant
    dashboard: DashboardAnalyticsResponse
    orders: list[Order]
    period_label: str
    generated_at: datetime
    logo: bytes | None
    zone: ZoneInfo
    brand: str = field(init=False)

    def __post_init__(self):
        self.brand = _hex(self.tenant.primary_color)

    @property
    def cancelled(self) -> list[Order]:
        return [o for o in self.orders if o.status is OrderStatus.CANCELLED]

    @property
    def discounts(self) -> Decimal:
        return sum((o.discount for o in self.orders if o.status is not OrderStatus.CANCELLED), Decimal("0"))

    @property
    def fees(self) -> Decimal:
        return sum((o.service_fee + o.delivery_fee for o in self.orders if o.status is not OrderStatus.CANCELLED), Decimal("0"))

    def local(self, moment: datetime) -> datetime:
        return moment.astimezone(self.zone).replace(tzinfo=None)


def build_sales_report(db: Session, tenant: Tenant, period: str, start: date | None = None, end: date | None = None) -> SalesReport:
    service = AnalyticsService(db, tenant)
    dashboard = service.dashboard(period, start, end, product_limit=None)
    zone = ZoneInfo(tenant.timezone)
    bounds_start = datetime.fromisoformat(dashboard.period_start) if dashboard.period_start else None
    bounds_end = datetime.fromisoformat(dashboard.period_end) if dashboard.period_end else None
    orders = OrderRepository(db, tenant.id).export(date_from=bounds_start, date_to=bounds_end, limit=20000)

    if bounds_start and bounds_end:
        last_day = (bounds_end - timedelta(days=1)).date()
        today = datetime.now(zone).date()
        shown_end = min(last_day, today) if bounds_start.date() <= today else last_day
        period_label = f"{bounds_start:%d/%m/%Y} a {shown_end:%d/%m/%Y}"
    elif orders:
        period_label = f"{orders[-1].created_at.astimezone(zone):%d/%m/%Y} a {datetime.now(zone):%d/%m/%Y} (todo o período)"
    else:
        period_label = "Todo o período"
    return SalesReport(
        tenant=tenant, dashboard=dashboard, orders=orders, period_label=period_label,
        generated_at=datetime.now(zone), logo=_logo_png(db, tenant), zone=zone,
    )


def weekday_summary(d: DashboardAnalyticsResponse) -> list[tuple[str, int, float, float]]:
    """(weekday, orders, revenue, average revenue per day that had sales), Monday first, days without sales left out."""
    totals: dict[int, list] = {}
    for item in d.sales_by_day:
        entry = totals.setdefault(date.fromisoformat(item.day).weekday(), [0, 0.0, 0])
        entry[0] += item.orders
        entry[1] += float(item.revenue)
        entry[2] += 1
    return [(WEEKDAYS[day], orders, value, value / days) for day, (orders, value, days) in sorted(totals.items())]


# ---- Excel ----------------------------------------------------------------------------------------

class _Sheet:
    """Small helper that keeps every sheet looking the same."""

    def __init__(self, sheet, brand: str):
        self.ws = sheet
        self.brand = brand
        self.header_fill = PatternFill("solid", fgColor=brand)
        self.soft_fill = PatternFill("solid", fgColor=_tint(brand, 0.9))
        self.zebra_fill = PatternFill("solid", fgColor=ZEBRA)
        self.border = Border(bottom=Side(style="thin", color=LINE))

    def title(self, report: SalesReport, subtitle: str, width: int) -> int:
        """Branded header on rows 1-3; returns the first free row."""
        ws = self.ws
        offset = 1
        if report.logo:
            image = XlImage(io.BytesIO(report.logo))
            image.width = image.height = 58
            ws.add_image(image, "A1")
            ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 10)
            offset = 2
        col = get_column_letter(offset)
        ws[f"{col}1"] = report.tenant.name
        ws[f"{col}1"].font = Font(size=18, bold=True, color=INK)
        ws[f"{col}2"] = f"{subtitle} · {report.period_label}"
        ws[f"{col}2"].font = Font(size=11, color=self.brand, bold=True)
        ws[f"{col}3"] = f"Gerado em {report.generated_at:%d/%m/%Y às %H:%M} · Mesa Digital"
        ws[f"{col}3"].font = Font(size=9, color=MUTED)
        for row in (1, 2, 3):
            ws.row_dimensions[row].height = 22 if row == 1 else 16
        for column in range(1, width + 1):
            ws.cell(row=4, column=column).border = Border(bottom=Side(style="medium", color=self.brand))
        return 6

    def section(self, row: int, text: str, column: int = 1) -> int:
        cell = self.ws.cell(row=row, column=column, value=text)
        cell.font = Font(size=12, bold=True, color=INK)
        return row + 1

    def table(self, row: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None,
              column: int = 1, total: list | None = None, muted_rows: set[int] | None = None) -> tuple[int, int]:
        """Header + zebra rows (+ total). Returns (first data row, next free row)."""
        formats = formats or {}
        ws = self.ws
        for index, text in enumerate(headers):
            cell = ws.cell(row=row, column=column + index, value=text)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal="right" if index in formats else "left", vertical="center")
        ws.row_dimensions[row].height = 20
        first = row + 1
        for offset, values in enumerate(rows):
            current = first + offset
            for index, value in enumerate(values):
                cell = ws.cell(row=current, column=column + index, value=safe_cell(value))
                if index in formats:
                    cell.number_format = formats[index]
                    cell.alignment = Alignment(horizontal="right")
                cell.border = self.border
                if offset % 2:
                    cell.fill = self.zebra_fill
                if muted_rows and offset in muted_rows:
                    cell.font = Font(color="9CA3AF", italic=True)
        after = first + len(rows)
        if total is not None:
            for index, value in enumerate(total):
                cell = ws.cell(row=after, column=column + index, value=value)
                cell.font = Font(bold=True, color=INK)
                cell.fill = self.soft_fill
                if index in formats and value not in (None, ""):
                    cell.number_format = formats[index]
                    cell.alignment = Alignment(horizontal="right")
            after += 1
        return first, after + 1

    def widths(self, values: dict[str, float]) -> None:
        for column, width in values.items():
            self.ws.column_dimensions[column].width = width

    def finish(self, header_row: int | None = None, landscape: bool = False) -> None:
        ws = self.ws
        ws.sheet_view.showGridLines = False
        if header_row:
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = "&8Mesa Digital"
        ws.oddFooter.right.text = "&8Página &P de &N"


def _bar_chart(title: str, y_title: str, money_axis: bool) -> BarChart:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.y_axis.title = y_title
    chart.y_axis.majorGridlines = None
    chart.legend = None
    chart.height, chart.width = 7.5, 22
    if money_axis:
        chart.y_axis.numFmt = '"R$" #,##0'
    return chart


def sales_report_xlsx(report: SalesReport) -> bytes:
    d = report.dashboard
    book = Workbook()
    revenue = float(d.revenue) or 0.0

    # Resumo
    summary = _Sheet(book.active, report.brand)
    summary.ws.title = "Resumo"
    row = summary.title(report, "Relatório de vendas", 6)
    cards = [
        ("Faturamento", float(d.revenue), MONEY_FORMAT),
        ("Pedidos", d.order_count, "0"),
        ("Ticket médio", float(d.average_ticket), MONEY_FORMAT),
        ("Vs. período anterior", float(d.revenue_change_percent) / 100 if d.previous_revenue else None, "+0.0%;-0.0%;0.0%"),
        ("Descontos dados", float(report.discounts), MONEY_FORMAT),
        ("Cancelados", len(report.cancelled), "0"),
    ]
    for index, (name, value, number_format) in enumerate(cards):
        head = summary.ws.cell(row=row, column=1 + index, value=name)
        head.font = Font(size=9, bold=True, color=MUTED)
        head.fill = summary.soft_fill
        head.alignment = Alignment(horizontal="left", indent=1)
        cell = summary.ws.cell(row=row + 1, column=1 + index, value=value if value is not None else "—")
        cell.font = Font(size=16, bold=True, color=report.brand if index == 0 else INK)
        cell.fill = summary.soft_fill
        cell.number_format = number_format
        cell.alignment = Alignment(horizontal="left", indent=1)
    summary.ws.row_dimensions[row + 1].height = 30
    row += 3

    def breakdown(title: str, items, mapping) -> list[list]:
        total = sum(i.count for i in items) or 1
        return [[label(mapping, i.label), i.count, i.count / total] for i in sorted(items, key=lambda i: -i.count)]

    row = summary.section(row, "Formas de pagamento")
    _, row = summary.table(row, ["Forma", "Pedidos", "% dos pedidos"], breakdown("", d.orders_by_payment, PAYMENT_LABELS), {1: "0", 2: "0.0%"})
    row = summary.section(row, "Tipo de atendimento")
    _, row = summary.table(row, ["Atendimento", "Pedidos", "% dos pedidos"], breakdown("", d.orders_by_service, SERVICE_LABELS), {1: "0", 2: "0.0%"})
    if d.categories:
        row = summary.section(row, "Faturamento por categoria")
        first, row = summary.table(
            row, ["Categoria", "Itens vendidos", "Faturamento", "% do faturamento"],
            [[c.category_name, c.quantity, float(c.revenue), percent(c.revenue, revenue)] for c in d.categories],
            {1: "0", 2: MONEY_FORMAT, 3: "0.0%"},
        )
        pie = PieChart()
        pie.title = "Faturamento por categoria"
        pie.add_data(Reference(summary.ws, min_col=3, min_row=first - 1, max_row=first + len(d.categories) - 1), titles_from_data=True)
        pie.set_categories(Reference(summary.ws, min_col=1, min_row=first, max_row=first + len(d.categories) - 1))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.height, pie.width = 8, 12
        summary.ws.add_chart(pie, f"F{first - 1}")
    summary.widths({"A": 26, "B": 16, "C": 16, "D": 20, "E": 18, "F": 14})
    summary.finish()

    # Vendas por dia
    daily = _Sheet(book.create_sheet("Vendas por dia"), report.brand)
    row = daily.title(report, "Vendas por dia", 5)
    days = [(date.fromisoformat(s.day), s.orders, float(s.revenue)) for s in d.sales_by_day]
    first, after = daily.table(
        row, ["Data", "Dia da semana", "Pedidos", "Faturamento", "Ticket médio"],
        [[day, WEEKDAYS[day.weekday()], orders, value, value / orders if orders else 0] for day, orders, value in days],
        {0: "dd/mm/yyyy", 2: "0", 3: MONEY_FORMAT, 4: MONEY_FORMAT},
        total=["Total", "", d.order_count, revenue, float(d.average_ticket)],
    )
    if days:
        chart = _bar_chart("Faturamento por dia", "R$", True)
        chart.add_data(Reference(daily.ws, min_col=4, min_row=first - 1, max_row=first + len(days) - 1), titles_from_data=True)
        chart.set_categories(Reference(daily.ws, min_col=1, min_row=first, max_row=first + len(days) - 1))
        chart.x_axis.number_format = "dd/mm"
        chart.series[0].graphicalProperties.solidFill = report.brand
        daily.ws.add_chart(chart, f"G{row}")
    weekdays = weekday_summary(d)
    if weekdays:
        row_w = daily.section(after + 1, "Por dia da semana")
        daily.table(row_w, ["Dia", "Pedidos", "Faturamento", "Média por dia aberto"],
                    [[name, orders, value, average] for name, orders, value, average in weekdays], {1: "0", 2: MONEY_FORMAT, 3: MONEY_FORMAT})
    daily.widths({"A": 14, "B": 16, "C": 11, "D": 16, "E": 15})
    daily.finish(header_row=row)

    # Horários
    hours = _Sheet(book.create_sheet("Horários"), report.brand)
    row = hours.title(report, "Movimento por horário", 4)
    by_hour = sorted(d.sales_by_hour, key=lambda h: h.hour)
    first, _ = hours.table(
        row, ["Horário", "Pedidos", "Faturamento", "% do faturamento"],
        [[f"{h.hour:02d}h às {(h.hour + 1) % 24:02d}h", h.orders, float(h.revenue), percent(h.revenue, revenue)] for h in by_hour],
        {1: "0", 2: MONEY_FORMAT, 3: "0.0%"},
    )
    if by_hour:
        chart = _bar_chart("Pedidos por horário", "Pedidos", False)
        chart.add_data(Reference(hours.ws, min_col=2, min_row=first - 1, max_row=first + len(by_hour) - 1), titles_from_data=True)
        chart.set_categories(Reference(hours.ws, min_col=1, min_row=first, max_row=first + len(by_hour) - 1))
        chart.series[0].graphicalProperties.solidFill = report.brand
        hours.ws.add_chart(chart, f"F{row}")
    hours.widths({"A": 16, "B": 11, "C": 16, "D": 18})
    hours.finish(header_row=row)

    # Produtos
    products = _Sheet(book.create_sheet("Produtos"), report.brand)
    row = products.title(report, "Ranking de produtos", 5)
    first, _ = products.table(
        row, ["#", "Produto", "Quantidade", "Faturamento", "% do faturamento"],
        [[i, p.product_name, p.quantity, float(p.revenue), percent(p.revenue, revenue)] for i, p in enumerate(d.products, start=1)],
        {0: "0", 2: "0", 3: MONEY_FORMAT, 4: "0.0%"},
    )
    if d.products:
        last = first + len(d.products) - 1
        products.ws.conditional_formatting.add(f"D{first}:D{last}", DataBarRule(start_type="min", end_type="max", color=_tint(report.brand, 0.35)))
        products.ws.auto_filter.ref = f"A{row}:E{last}"
    products.widths({"A": 6, "B": 38, "C": 12, "D": 16, "E": 18})
    products.finish(header_row=row)

    # Pedidos
    _orders_sheet(_Sheet(book.create_sheet("Pedidos"), report.brand), report)

    for sheet in book.worksheets:
        neutralise_formulas(sheet)
    book.properties.title = f"{report.tenant.name} - Relatório de vendas"
    book.properties.creator = "Mesa Digital"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


_ORDER_HEADERS = ["Nº", "Data", "Hora", "Cliente", "Telefone", "Atendimento", "Mesa", "Status", "Pagamento", "Forma",
                  "Itens", "Subtotal", "Desconto", "Taxa de serviço", "Entrega", "Total"]
_ORDER_FORMATS = {0: "0", 1: "dd/mm/yyyy", 2: "hh:mm", 11: MONEY_FORMAT, 12: MONEY_FORMAT, 13: MONEY_FORMAT, 14: MONEY_FORMAT, 15: MONEY_FORMAT}


def _order_rows(orders: list[Order], zone: ZoneInfo) -> tuple[list[list], set[int]]:
    rows, cancelled = [], set()
    for index, order in enumerate(orders):
        moment = order.created_at.astimezone(zone).replace(tzinfo=None)
        if order.status is OrderStatus.CANCELLED:
            cancelled.add(index)
        rows.append([
            order.order_number, moment.date(), moment.time().replace(second=0, microsecond=0),
            order.customer_name or "", order.customer_phone or "", label(SERVICE_LABELS, order.service_type),
            order.table_label or "", label(STATUS_LABELS, order.status),
            "Pago" if order.payment_status is PaymentStatus.PAID else "Pendente", label(PAYMENT_LABELS, order.payment_method),
            ", ".join(f"{i.quantity}x {i.product_name}" for i in order.items),
            float(order.subtotal), float(order.discount), float(order.service_fee), float(order.delivery_fee), float(order.total),
        ])
    return rows, cancelled


def _orders_sheet(sheet: _Sheet, report: SalesReport, subtitle: str = "Pedidos do período") -> None:
    row = sheet.title(report, subtitle, len(_ORDER_HEADERS))
    rows, cancelled = _order_rows(report.orders, report.zone)
    valid = [o for o in report.orders if o.status is not OrderStatus.CANCELLED]
    total = ["Total (sem cancelados)", "", "", f"{len(valid)} pedidos", "", "", "", "", "", "", "",
             float(sum((o.subtotal for o in valid), Decimal(0))), float(sum((o.discount for o in valid), Decimal(0))),
             float(sum((o.service_fee for o in valid), Decimal(0))), float(sum((o.delivery_fee for o in valid), Decimal(0))),
             float(sum((o.total for o in valid), Decimal(0)))]
    first, _ = sheet.table(row, _ORDER_HEADERS, rows, _ORDER_FORMATS, total=total if rows else None, muted_rows=cancelled)
    if rows:
        sheet.ws.auto_filter.ref = f"A{row}:{get_column_letter(len(_ORDER_HEADERS))}{first + len(rows) - 1}"
    sheet.widths({"A": 7, "B": 12, "C": 8, "D": 24, "E": 16, "F": 13, "G": 8, "H": 12, "I": 11, "J": 12,
                  "K": 48, "L": 12, "M": 11, "N": 14, "O": 11, "P": 12})
    sheet.finish(header_row=row, landscape=True)


def orders_xlsx(db: Session, tenant: Tenant, orders: list[Order], period_label: str) -> bytes:
    zone = ZoneInfo(tenant.timezone)
    report = SalesReport(tenant=tenant, dashboard=None, orders=orders, period_label=period_label,  # type: ignore[arg-type]
                         generated_at=datetime.now(zone), logo=_logo_png(db, tenant), zone=zone)
    book = Workbook()
    _orders_sheet(_Sheet(book.active, report.brand), report, "Histórico de pedidos")
    book.active.title = "Pedidos"
    neutralise_formulas(book.active)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def customers_xlsx(db: Session, tenant: Tenant, customers: list) -> bytes:
    zone = ZoneInfo(tenant.timezone)
    report = SalesReport(tenant=tenant, dashboard=None, orders=[], period_label=f"{len(customers)} clientes",  # type: ignore[arg-type]
                         generated_at=datetime.now(zone), logo=_logo_png(db, tenant), zone=zone)
    book = Workbook()
    sheet = _Sheet(book.active, report.brand)
    sheet.ws.title = "Clientes"
    row = sheet.title(report, "Clientes", 6)
    rows = [[c.name, c.phone or "", c.orders_count, float(c.total_spent),
             c.last_order_at.astimezone(zone).replace(tzinfo=None).date() if c.last_order_at else None, c.notes or ""] for c in customers]
    first, _ = sheet.table(row, ["Nome", "Telefone", "Pedidos", "Total gasto", "Último pedido", "Observações"], rows,
                           {2: "0", 3: MONEY_FORMAT, 4: "dd/mm/yyyy"})
    if rows:
        sheet.ws.auto_filter.ref = f"A{row}:F{first + len(rows) - 1}"
    sheet.widths({"A": 30, "B": 17, "C": 10, "D": 14, "E": 14, "F": 44})
    sheet.finish(header_row=row)
    neutralise_formulas(sheet.ws)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


# ---- PDF ----------------------------------------------------------------------------------------

class _NumberedCanvas(pdf_canvas.Canvas):
    """Draws "Página X de Y" once the total page count is known."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for state in self._pages:
            self.__dict__.update(state)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#" + MUTED))
            self.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Página {self._pageNumber} de {total}")
            super().showPage()
        super().save()


def sales_report_pdf(report: SalesReport) -> bytes:
    d = report.dashboard
    brand = colors.HexColor("#" + report.brand)
    soft = colors.HexColor("#" + _tint(report.brand, 0.92))
    ink, muted, line = colors.HexColor("#" + INK), colors.HexColor("#" + MUTED), colors.HexColor("#" + LINE)
    revenue = Decimal(d.revenue or 0)
    width = A4[0] - 36 * mm

    styles = {
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, textColor=ink, spaceBefore=10, spaceAfter=6, keepWithNext=1),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=8, textColor=muted, leading=10),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=9, textColor=ink, leading=11),
        "cell_r": ParagraphStyle("cell_r", fontName="Helvetica", fontSize=9, textColor=ink, leading=11, alignment=TA_RIGHT),
    }

    def header_footer(canvas, doc):
        canvas.saveState()
        top = A4[1]
        canvas.setFillColor(brand)
        canvas.rect(0, top - 30 * mm, A4[0], 30 * mm, stroke=0, fill=1)
        x = 18 * mm
        if report.logo:
            canvas.setFillColor(colors.white)
            canvas.roundRect(x, top - 25 * mm, 20 * mm, 20 * mm, 3 * mm, stroke=0, fill=1)
            canvas.drawImage(ImageReader(io.BytesIO(report.logo)), x + 1.5 * mm, top - 23.5 * mm, 17 * mm, 17 * mm, mask="auto", preserveAspectRatio=True)
            x += 25 * mm
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 17)
        canvas.drawString(x, top - 14 * mm, report.tenant.name[:48])
        canvas.setFont("Helvetica", 10)
        canvas.drawString(x, top - 20.5 * mm, f"Relatório de vendas · {report.period_label}")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0] - 18 * mm, top - 14 * mm, f"Gerado em {report.generated_at:%d/%m/%Y %H:%M}")
        canvas.setFillColor(muted)
        canvas.drawString(18 * mm, 10 * mm, "Mesa Digital · mesadig.com.br")
        canvas.restoreState()

    def kpis() -> Table:
        change = "—"
        if d.previous_revenue:
            value = float(d.revenue_change_percent)
            change = f"{'+' if value >= 0 else ''}{value:.1f}%".replace(".", ",")
        items = [("Faturamento", money(d.revenue)), ("Pedidos", str(d.order_count)), ("Ticket médio", money(d.average_ticket)),
                 ("Vs. período anterior", change)]
        cells = [[Paragraph(f'<font size="8" color="#{MUTED}">{name.upper()}</font><br/><font size="15"><b>{value}</b></font>', ParagraphStyle("k", leading=20))
                  for name, value in items]]
        table = Table(cells, colWidths=[width / 4 - 2 * mm] * 4, rowHeights=[17 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), soft), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm), ("LINEBEFORE", (1, 0), (-1, 0), 2, colors.white),
            ("LINEABOVE", (0, 0), (-1, 0), 2, brand),
        ]))
        return table

    def data_table(headers: list[str], rows: list[list[str]], widths: list[float], right: set[int], total: list[str] | None = None) -> Table:
        body = [[Paragraph(f"<b>{h}</b>", ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.white,
                                                         alignment=TA_RIGHT if i in right else 0)) for i, h in enumerate(headers)]]
        for values in rows:
            body.append([Paragraph(str(v), styles["cell_r"] if i in right else styles["cell"]) for i, v in enumerate(values)])
        if total:
            body.append([Paragraph(f"<b>{v}</b>", styles["cell_r"] if i in right else styles["cell"]) for i, v in enumerate(total)])
        table = Table(body, colWidths=widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), brand), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LINEBELOW", (0, 1), (-1, -1), 0.4, line),
        ]
        for index in range(2, len(body), 2):
            style.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#" + ZEBRA)))
        if total:
            style.append(("BACKGROUND", (0, -1), (-1, -1), soft))
        table.setStyle(TableStyle(style))
        return table

    def daily_chart() -> Drawing | None:
        days = d.sales_by_day
        if not days:
            return None
        drawing = Drawing(width, 62 * mm)
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = 14 * mm, 10 * mm, width - 18 * mm, 46 * mm
        chart.data = [[float(s.revenue) for s in days]]
        step = max(1, -(-len(days) // 12))
        chart.categoryAxis.categoryNames = [f"{date.fromisoformat(s.day):%d/%m}" if i % step == 0 else "" for i, s in enumerate(days)]
        chart.categoryAxis.labels.fontSize = 7
        chart.categoryAxis.labels.fontName = "Helvetica"
        chart.valueAxis.labels.fontName = "Helvetica"
        chart.categoryAxis.labels.fillColor = muted
        chart.categoryAxis.strokeColor = line
        chart.valueAxis.labels.fontSize = 7
        chart.valueAxis.labels.fillColor = muted
        chart.valueAxis.strokeColor = colors.white
        chart.valueAxis.gridStrokeColor = line
        chart.valueAxis.visibleGrid = True
        chart.valueAxis.valueMin = 0
        chart.valueAxis.labelTextFormat = lambda v: f"R$ {v:,.0f}".replace(",", ".")
        chart.bars[0].fillColor = brand
        chart.bars[0].strokeColor = None
        chart.barSpacing = 1
        drawing.add(chart)
        drawing.add(String(0, 58 * mm, "Faturamento por dia", fontName="Helvetica-Bold", fontSize=11, fillColor=ink))
        return drawing

    story: list = [kpis(), Spacer(1, 4 * mm)]
    extra = f"Descontos concedidos: {money(report.discounts)} · Taxas (serviço e entrega): {money(report.fees)} · Pedidos cancelados: {len(report.cancelled)} (fora dos totais)"
    story += [Paragraph(extra, styles["small"]), Spacer(1, 5 * mm)]
    chart = daily_chart()
    if chart:
        story += [chart, Spacer(1, 2 * mm)]

    def breakdown_rows(items, mapping) -> list[list[str]]:
        total = sum(i.count for i in items) or 1
        return [[label(mapping, i.label), str(i.count), f"{i.count / total:.0%}"] for i in sorted(items, key=lambda i: -i.count)]

    half = width / 2 - 3 * mm
    pay = data_table(["Forma de pagamento", "Pedidos", "%"], breakdown_rows(d.orders_by_payment, PAYMENT_LABELS), [half * 0.56, half * 0.24, half * 0.2], {1, 2})
    svc = data_table(["Atendimento", "Pedidos", "%"], breakdown_rows(d.orders_by_service, SERVICE_LABELS), [half * 0.56, half * 0.24, half * 0.2], {1, 2})
    pair = Table([[pay, svc]], colWidths=[width / 2, width / 2])
    pair.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (0, 0), 6 * mm)]))
    story += [KeepTogether([Paragraph("Como os clientes compram", styles["h2"]), pair])]

    if d.sales_by_hour:
        busiest = sorted(d.sales_by_hour, key=lambda h: -h.orders)[:8]
        story += [KeepTogether([
            Paragraph("Horários de maior movimento", styles["h2"]),
            data_table(["Horário", "Pedidos", "Faturamento"],
                       [[f"{h.hour:02d}h às {(h.hour + 1) % 24:02d}h", str(h.orders), money(h.revenue)] for h in busiest],
                       [width - 60 * mm, 25 * mm, 35 * mm], {1, 2}),
        ])]
    if d.products:
        top = d.products[:15]
        story += [Paragraph("Produtos mais vendidos", styles["h2"]), (
            data_table(["#", "Produto", "Qtd.", "Faturamento", "% do total"],
                       [[str(i), p.product_name, str(p.quantity), money(p.revenue), f"{percent(p.revenue, revenue):.1%}".replace(".", ",")]
                        for i, p in enumerate(top, start=1)],
                       [8 * mm, width - 78 * mm, 18 * mm, 28 * mm, 24 * mm], {0, 2, 3, 4}))]
    if d.categories:
        story += [KeepTogether([
            Paragraph("Faturamento por categoria", styles["h2"]),
            data_table(["Categoria", "Itens", "Faturamento", "% do total"],
                       [[c.category_name, str(c.quantity), money(c.revenue), f"{percent(c.revenue, revenue):.1%}".replace(".", ",")] for c in d.categories],
                       [width - 70 * mm, 18 * mm, 28 * mm, 24 * mm], {1, 2, 3},
                       total=["Total", str(sum(c.quantity for c in d.categories)), money(revenue), "100%"]),
        ])]
    weekdays = weekday_summary(d)
    if weekdays:
        story += [KeepTogether([
            Paragraph("Movimento por dia da semana", styles["h2"]),
            data_table(["Dia", "Pedidos", "Faturamento", "Média por dia aberto"],
                       [[name.capitalize(), str(orders), money(value), money(average)] for name, orders, value, average in weekdays],
                       [width - 96 * mm, 22 * mm, 34 * mm, 40 * mm], {1, 2, 3}),
        ])]
    if not d.order_count:
        story.append(Paragraph("Nenhuma venda neste período.", styles["cell"]))

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=38 * mm, bottomMargin=18 * mm,
                                 title=f"{report.tenant.name} - Relatório de vendas", author="Mesa Digital")
    document.build(story, onFirstPage=header_footer, onLaterPages=header_footer, canvasmaker=_NumberedCanvas)
    return output.getvalue()
